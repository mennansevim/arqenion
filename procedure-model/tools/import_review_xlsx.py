# -*- coding: utf-8 -*-
"""Doldurulmuş onay Excel'i → review_decisions.json (clinical_review.py'ye işlenecek kararlar)."""
import json, sys
from openpyxl import load_workbook
wb = load_workbook(sys.argv[1], data_only=True)
def rows(name, id_col=1, dec_col=None, note_col=None):
    out = []
    for r in wb[name].iter_rows(min_row=2, values_only=True):
        if not r[id_col - 1]: continue
        dec = (r[dec_col - 1] or "").strip() if dec_col else None
        note = (r[note_col - 1] or "").strip() if note_col else ""
        if dec or note: out.append({"id": r[id_col - 1], "decision": dec, "note": note})
    return out
res = {"drafts": rows("Taslak senaryolar", 1, 11, 12),
       "existing": rows("Mevcut senaryolar", 1, 6, 7),
       "questions": [{"id": r[0], "answer": str(r[3]).strip()} for r in wb["Bekleyen sorular"].iter_rows(min_row=2, values_only=True) if r[0] and r[3]]}
json.dump(res, open(sys.argv[2] if len(sys.argv) > 2 else "review_decisions.json", "w"), ensure_ascii=False, indent=1)
print({k: len(v) for k, v in res.items()})
pending = [r[0] for r in wb["Taslak senaryolar"].iter_rows(min_row=2, values_only=True) if r[0] and not r[10]]
if pending: print("Karar verilmemiş taslaklar:", pending)
