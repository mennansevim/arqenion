# -*- coding: utf-8 -*-
"""Model → klinik onay Excel'i (taslak senaryolar, mevcut senaryolar, bekleyen sorular)."""
import json, sys
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

m = json.load(open(sys.argv[1])); OUT = sys.argv[2]
S = {s["step"]: s for s in m["steps"]}
F = "Arial"
HEAD = PatternFill("solid", fgColor="1F3A4D"); FILL = PatternFill("solid", fgColor="FFF2B3"); GREY = PatternFill("solid", fgColor="F2F2F2")
thin = Side(style="thin", color="BFBFBF"); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")
STAT = {"ACTIVE": "Aynen kalıyor", "ADAPTED": "Metni yeni karara göre uyarlandı", "SUPERSEDED": "Devreden çıktı", "CONDITIONAL": "Koşullu"}

def bullets(xs): return "\n".join(f"• {x}" for x in xs)
def numbered(xs): return "\n".join(f"{i + 1}. {x}" for i, x in enumerate(xs))

def table(ws, headers, widths, rows, input_cols, row_h=None):
    for j, (h, w) in enumerate(zip(headers, widths), 1):
        c = ws.cell(row=1, column=j, value=h)
        c.font = Font(name=F, bold=True, color="FFFFFF", size=11); c.fill = FILL if j in input_cols else HEAD
        if j in input_cols: c.font = Font(name=F, bold=True, color="000000", size=11)
        c.alignment = Alignment(wrap_text=True, vertical="center"); c.border = BOX
        ws.column_dimensions[c.column_letter].width = w
    for i, r in enumerate(rows, 2):
        for j, v in enumerate(r, 1):
            c = ws.cell(row=i, column=j, value=v)
            c.font = Font(name=F, size=10); c.alignment = WRAP; c.border = BOX
            if j in input_cols: c.fill = FILL
        if row_h:
            import math
            lines = max(sum(max(1, math.ceil(len(seg) / (widths[j - 1] * 1.15))) for seg in str(v or "").split("\n"))
                        for j, v in enumerate(r, 1))
            ws.row_dimensions[i].height = max(row_h, 13 * lines + 8)
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "D2"

def dropdown(ws, col, n, options):
    dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True,
                        promptTitle="Karar", prompt="Listeden seçin", showInputMessage=True)
    ws.add_data_validation(dv); dv.add(f"{col}2:{col}{n + 1}")

wb = Workbook()
# ---------- 1. Nasıl doldurulur
ws = wb.active; ws.title = "Nasıl doldurulur"
lines = [
    ("Açık Apendektomi: Hata Senaryoları Onayı (model v" + m["procedure"]["version"] + ")", True),
    ("", False),
    ("Yalnızca SARI sütunları doldurun. Diğer sütunlar bilgi içindir, değiştirmeyin.", True),
    ("", False),
    ("1. 'Taslak senaryolar' sayfası: Adım 50, 52 ve 56 için yazılan 6 yeni senaryo. Her satırda KARAR sütunundan seçin:", False),
    ("     Onay = metin doğru, modele girsin", False),
    ("     Düzeltme = NOT sütununa neyin değişeceğini yazın", False),
    ("     Gerek yok = bu senaryo istenmiyor (nedenini NOT'a kısaca yazın)", False),
    ("2. 'Mevcut senaryolar' sayfası: Bu üç adımdaki eski senaryoların durumu. Uygun değilse KARAR'dan 'Kaldırılsın' ya da 'Düzeltme' seçin.", False),
    ("3. 'Bekleyen sorular' sayfası: Geçen turdan kalan 5 kısa soru. CEVAP sütununa yazın.", False),
    ("", False),
    ("Örnek doldurma (Taslak senaryolar sayfasında bir satır):", True),
]
for i, (t, b) in enumerate(lines, 1):
    c = ws.cell(row=i, column=1, value=t); c.font = Font(name=F, size=12 if i == 1 else 11, bold=b)
ex = [("ID", "KARAR", "NOT"),
      ("S50-D1", "Onay", ""),
      ("S50-D2", "Gerek yok", "Vicryl Rapide bizim ameliyathanede yok, karıştırılma ihtimali yok."),
      ("S56-D1", "Düzeltme", "Belirtilerden 'genellikle mavi renkli' kısmı çıksın. Çözümün 3. maddesine 'gerekirse Metzenbaum makas uzatılır' eklensin.")]
r0 = len(lines) + 1
for i, row in enumerate(ex):
    for j, v in enumerate(row, 1):
        c = ws.cell(row=r0 + i, column=j, value=v); c.border = BOX; c.alignment = WRAP
        c.font = Font(name=F, size=10, bold=(i == 0), italic=(i > 0), color="595959" if i > 0 else "000000")
        if i == 0: c.fill = GREY
        elif j > 1: c.fill = FILL
ws.cell(row=r0 + len(ex) + 1, column=1, value="(Bu tablo yalnızca örnektir; gerçek kararlar 'Taslak senaryolar' sayfasına yazılır.)").font = Font(name=F, size=9, italic=True, color="7F7F7F")
ws.column_dimensions["A"].width = 14; ws.column_dimensions["B"].width = 14; ws.column_dimensions["C"].width = 90

# ---------- 2. Taslak senaryolar
ws = wb.create_sheet("Taslak senaryolar")
rows = []
for n in (50, 52, 56):
    for e in S[n]["scrub"]["errorScenarios"]:
        if e["reviewStatus"] != "DRAFT": continue
        rows.append([e["id"], n, e["title"], e["review"].get("replaces") or "yeni", e["scrubError"],
                     bullets(e["possibleComplications"]), bullets(e["signs"]), numbered(e["resolutionSteps"]),
                     e["caseContinuation"], bullets(e["organizationalImpact"]), "", ""])
table(ws, ["ID", "Adım", "Senaryo başlığı", "Yerini aldığı", "Scrub'ın yaptığı yanlışlık", "Olası komplikasyon",
           "Komplikasyon belirtileri", "Scrub'ın sıralı çözüm yaklaşımı", "Vakanın devamı", "Yönetim ve organizasyona etkisi",
           "KARAR", "NOT"],
      [9, 6, 32, 11, 40, 50, 42, 60, 36, 40, 14, 45], rows, input_cols={11, 12}, row_h=60)
dropdown(ws, "K", len(rows), ["Onay", "Düzeltme", "Gerek yok"])

# ---------- 3. Mevcut senaryolar
ws = wb.create_sheet("Mevcut senaryolar")
rows = []
for n in (50, 52, 56):
    for e in S[n]["scrub"]["errorScenarios"]:
        if e["reviewStatus"] == "DRAFT": continue
        rows.append([e["id"], n, e["title"], STAT[e["reviewStatus"]], e.get("review", {}).get("note", ""), "", ""])
table(ws, ["ID", "Adım", "Senaryo başlığı", "Durum", "Açıklama", "KARAR", "NOT"], [9, 6, 55, 30, 45, 14, 45],
      rows, input_cols={6, 7}, row_h=30)
dropdown(ws, "F", len(rows), ["Uygun", "Kaldırılsın", "Düzeltme"])

# ---------- 4. Bekleyen sorular
ws = wb.create_sheet("Bekleyen sorular")
rows = [[q["id"], q.get("step", ""), q["question"], ""] for q in m["openQuestions"] if q["id"].startswith("Q-B")]
table(ws, ["ID", "Adım", "Soru", "CEVAP"], [8, 6, 90, 50], rows, input_cols={4}, row_h=40)
for w in wb.worksheets:
    w.page_setup.orientation = "landscape"; w.page_setup.paperSize = w.PAPERSIZE_A4
    w.page_setup.fitToWidth = 1; w.page_setup.fitToHeight = 0; w.sheet_properties.pageSetUpPr.fitToPage = True
    w.print_title_rows = "1:1" if w.title != "Nasıl doldurulur" else None
wb.save(OUT)
print(OUT, {w.title: w.max_row - 1 for w in wb.worksheets[1:]})
